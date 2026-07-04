"""
Adoption views (View layer).
Delegates to services (Presenter layer).
"""

from django.http import HttpResponse
from docx import Document
from rest_framework import permissions, status, viewsets
from rest_framework.decorators import action
from rest_framework.response import Response

from .models import Adoption
from .permissions import IsAdminOrCenterAdmin, IsApplicantOrCenterAdmin
from .serializers import (
    AdoptionCreateSerializer,
    AdoptionSerializer,
    AdoptionTimelineSerializer,
)
from .services import AdoptionService


class AdoptionViewSet(viewsets.ModelViewSet):
    """
    Manage adoption requests.

    Authenticated users (applicants):
        - POST /adoptions/ (create request)
        - GET /adoptions/ (own requests)
        - GET /adoptions/{id}/ (own request)
        - POST /adoptions/{id}/submit/
        - GET /adoptions/{id}/timeline/

    CenterAdmin / SuperAdmin only:
        - POST /adoptions/{id}/start_review/
        - POST /adoptions/{id}/approve/
        - POST /adoptions/{id}/reject/
        - POST /adoptions/{id}/complete/
    """

    queryset = Adoption.objects.select_related("pet", "applicant", "center", "reviewed_by")
    serializer_class = AdoptionSerializer

    def get_permissions(self):
        admin_actions = {"start_review", "approve", "reject", "complete"}
        if self.action in admin_actions:
            return [permissions.IsAuthenticated(), IsAdminOrCenterAdmin()]
        return [permissions.IsAuthenticated(), IsApplicantOrCenterAdmin()]

    def get_queryset(self):
        user = self.request.user
        qs = Adoption.objects.select_related("pet", "applicant", "center", "reviewed_by")
        pet_id = self.request.query_params.get("pet")
        if pet_id:
            qs = qs.filter(pet_id=pet_id)
        if self.request.query_params.get("mine") == "true":
            return qs.filter(applicant=user)
        if user.is_superuser:
            return qs.all()
        elif user.role == "center_admin":
            return qs.filter(center=user.center)
        return qs.filter(applicant=user)

    def get_serializer_class(self):
        if self.action == "create":
            return AdoptionCreateSerializer
        return AdoptionSerializer

    def get_serializer_context(self):
        context = super().get_serializer_context()
        context["request"] = self.request
        return context

    @action(
        detail=False,
        methods=["get"],
        permission_classes=[permissions.IsAuthenticated, IsAdminOrCenterAdmin],
    )
    def export(self, request):
        """Export adoptions to Word document."""
        qs = self.get_queryset()
        doc = Document()
        doc.add_heading("Reporte de Solicitudes de Adopción", 0)

        table = doc.add_table(rows=1, cols=6)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        for i, text in enumerate(["ID", "Solicitante", "Mascota", "Centro", "Estado", "Fecha"]):
            hdr[i].text = text

        for adoption in qs:
            row = table.add_row().cells
            row[0].text = str(adoption.id)
            row[1].text = adoption.applicant.get_full_name() or adoption.applicant.email
            row[2].text = adoption.pet.name if adoption.pet else ""
            row[3].text = adoption.center.name if adoption.center else ""
            row[4].text = adoption.get_status_display()
            row[5].text = adoption.created_at.strftime("%d/%m/%Y")

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        response["Content-Disposition"] = 'attachment; filename="solicitudes.docx"'
        doc.save(response)
        return response

    def perform_create(self, serializer):
        adoption = serializer.save(applicant=self.request.user)
        service = AdoptionService(adoption)
        service.submit()

    @action(detail=True, methods=["post"])
    def submit(self, request, pk=None):
        """Submit an adoption request."""
        adoption = self.get_object()
        service = AdoptionService(adoption)

        try:
            service.submit()
            return Response(AdoptionSerializer(adoption).data)
        except Exception as e:
            return Response({"error": str(e)}, status=status.HTTP_400_BAD_REQUEST)

    @action(
        detail=True,
        methods=["post"],
        permission_classes=[IsAdminOrCenterAdmin],
    )
    def start_review(self, request, pk=None):
        """Start reviewing an adoption request (admin only)."""
        adoption = self.get_object()
        service = AdoptionService(adoption)

        try:
            service.start_review(reviewed_by=request.user)
            return Response(AdoptionSerializer(adoption).data)
        except Exception as e:
            return Response({"error": str(e)}, status=status.HTTP_400_BAD_REQUEST)

    @action(
        detail=True,
        methods=["post"],
        permission_classes=[IsAdminOrCenterAdmin],
    )
    def approve(self, request, pk=None):
        """Approve an adoption request (admin only)."""
        adoption = self.get_object()
        service = AdoptionService(adoption)
        notes = request.data.get("notes", "")

        try:
            service.approve(approved_by=request.user, notes=notes)
            return Response(AdoptionSerializer(adoption).data)
        except Exception as e:
            return Response({"error": str(e)}, status=status.HTTP_400_BAD_REQUEST)

    @action(
        detail=True,
        methods=["post"],
        permission_classes=[IsAdminOrCenterAdmin],
    )
    def reject(self, request, pk=None):
        """Reject an adoption request (admin only)."""
        adoption = self.get_object()
        service = AdoptionService(adoption)
        reason = request.data.get("reason", "")

        try:
            service.reject(rejected_by=request.user, reason=reason)
            return Response(AdoptionSerializer(adoption).data)
        except Exception as e:
            return Response({"error": str(e)}, status=status.HTTP_400_BAD_REQUEST)

    @action(
        detail=True,
        methods=["post"],
        permission_classes=[IsAdminOrCenterAdmin],
    )
    def complete(self, request, pk=None):
        """Mark adoption as completed (admin only)."""
        adoption = self.get_object()
        service = AdoptionService(adoption)

        try:
            service.complete(completed_by=request.user)
            return Response(AdoptionSerializer(adoption).data)
        except Exception as e:
            return Response({"error": str(e)}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=["get"])
    def timeline(self, request, pk=None):
        """Get adoption timeline."""
        adoption = self.get_object()
        service = AdoptionService(adoption)
        timeline = service.get_timeline()

        return Response(AdoptionTimelineSerializer(timeline, many=True).data)
