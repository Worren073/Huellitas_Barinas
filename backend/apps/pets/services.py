"""
Pet services (Presenter layer).
Handles business logic for pet operations.
"""

from django.http import HttpResponse
from docx import Document

from .models import Pet, PetImage


class PetService:
    """Service class for pet operations.

    Usage:
        PetService(pet).mark_as_adopted()
        PetService.get_available_pets()  # classmethod
    """

    def __init__(self, pet):
        self.pet = pet

    def update_pet(self, **kwargs):
        """Update pet with validation."""
        for key, value in kwargs.items():
            setattr(self.pet, key, value)
        self.pet.full_clean()
        self.pet.save()
        return self.pet

    def mark_as_adopted(self):
        """Mark a pet as adopted."""
        self.pet.status = Pet.Status.ADOPTED
        self.pet.save()
        return self.pet

    def mark_as_in_process(self):
        """Mark a pet as in adoption process."""
        self.pet.status = Pet.Status.IN_PROCESS
        self.pet.save()
        return self.pet

    def mark_as_available(self):
        """Mark a pet as available for adoption."""
        self.pet.status = Pet.Status.AVAILABLE
        self.pet.save()
        return self.pet

    def mark_as_not_available(self):
        """Mark a pet as not available."""
        self.pet.status = Pet.Status.NOT_AVAILABLE
        self.pet.save()
        return self.pet

    @classmethod
    def create_pet(cls, **kwargs):
        """Create a new pet."""
        pet = Pet(**kwargs)
        pet.full_clean()
        pet.save()
        return cls(pet)

    @classmethod
    def get_pets_by_center(cls, center):
        """Get all pets for a specific center."""
        return Pet.objects.filter(center=center)

    @classmethod
    def get_available_pets(cls):
        """Get all available pets."""
        return Pet.objects.filter(status=Pet.Status.AVAILABLE)

    @classmethod
    def get_stats(cls, user):
        """Get dashboard stats scoped by user role."""
        base_qs = Pet.objects
        if user.role == "center_admin" and user.center:
            base_qs = base_qs.filter(center=user.center)
        return {
            "pets_count": base_qs.count(),
            "available_pets": base_qs.filter(status="available").count(),
            "in_process_pets": base_qs.filter(status="in_process").count(),
            "adopted_pets": base_qs.filter(status="adopted").count(),
        }

    @classmethod
    def export_to_docx(cls, queryset):
        """Generate a Word document with pets data."""
        doc = Document()
        doc.add_heading("Reporte de Mascotas", 0)

        table = doc.add_table(rows=1, cols=7)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        headers = ["Nombre", "Especie", "Raza", "Edad", "Tamaño", "Estado", "Centro"]
        for i, text in enumerate(headers):
            hdr[i].text = text

        for pet in queryset:
            row = table.add_row().cells
            row[0].text = pet.name
            row[1].text = pet.get_species_display()
            row[2].text = pet.breed
            row[3].text = f"{pet.age_months} meses"
            row[4].text = pet.get_size_display()
            row[5].text = pet.get_status_display()
            row[6].text = pet.center.name if pet.center else ""

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        response["Content-Disposition"] = 'attachment; filename="mascotas.docx"'
        doc.save(response)
        return response


class PetImageService:
    """Service class for pet image operations."""

    @staticmethod
    def create_image(pet_id, validated_data):
        """Create a pet image linked to a pet. WebP conversion via signal."""
        pet_image = PetImage(pet_id=pet_id, **validated_data)
        pet_image.full_clean()
        pet_image.save()
        return pet_image
